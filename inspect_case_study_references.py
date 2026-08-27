import docx
import fitz  # PyMuPDF
import sys

sys.stdout.reconfigure(encoding='utf-8')

print("=== Reading 學習共同體公開課的實踐歷程-周尚枚.docx ===")
try:
    doc1 = docx.Document('學習共同體公開課的實踐歷程-周尚枚.docx')
    for p in doc1.paragraphs:
        if p.text.strip():
            print(p.text[:200])
    for t in doc1.tables:
        for r in t.rows:
            print([c.text.strip().replace('\n', ' ')[:50] for c in r.cells])
except Exception as e:
    print(f"Docx error: {e}")

print("\n=== Reading 學生主體性彰顯的學共課堂_內觀人員的課例研究_周尚枚.pdf ===")
try:
    pdf_doc = fitz.open('學生主體性彰顯的學共課堂_內觀人員的課例研究_周尚枚.pdf')
    print(f"Total PDF Pages: {len(pdf_doc)}")
    for i in range(min(5, len(pdf_doc))):
        print(f"--- Page {i+1} ---")
        print(pdf_doc[i].get_text("text")[:500])
except Exception as e:
    print(f"PDF error: {e}")
