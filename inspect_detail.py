import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc1 = docx.Document('114教學省思心得紀錄表.docx')
print("=== doc1 Paragraphs ===")
for i, p in enumerate(doc1.paragraphs):
    print(f"P{i}: {p.text}")

print("=== doc1 Table 1 ===")
t1 = doc1.tables[0]
for r_idx, row in enumerate(t1.rows):
    print(f"Row {r_idx} (len {len(row.cells)}): {[c.text.strip().replace('\n', ' ') for c in row.cells]}")

doc2 = docx.Document('114觀課紀錄表.docx')
print("\n=== doc2 Paragraphs ===")
for i, p in enumerate(doc2.paragraphs):
    print(f"P{i}: {p.text}")

print("=== doc2 Table 1 ===")
t2 = doc2.tables[0]
for r_idx, row in enumerate(t2.rows):
    print(f"Row {r_idx} (len {len(row.cells)}): {[c.text.strip().replace('\n', ' ') for c in row.cells]}")
