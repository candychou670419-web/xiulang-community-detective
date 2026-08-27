import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import qn, nsdecls
import os

os.makedirs('output', exist_ok=True)
doc = docx.Document()

# Set margins to 0.75 inch
for section in doc.sections:
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/></w:tblBorders>')
    tblPr.append(borders)

# Title
p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_title = p_title.add_run('國小國語科 國字與注音學習對照表')
run_title.font.name = '標楷體'
run_title._element.rPr.rFonts.set(qn('w:eastAsia'), '標楷體')
run_title.font.size = Pt(18)
run_title.font.bold = True
run_title.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

# Subtitle / Header info
p_info = doc.add_paragraph()
p_info.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run_info = p_info.add_run('___年 ___班  座號：___  姓名：___________  得分：_______')
run_info.font.name = '微軟正黑體'
run_info._element.rPr.rFonts.set(qn('w:eastAsia'), '微軟正黑體')
run_info.font.size = Pt(11)

p_h1 = doc.add_paragraph()
r_h1 = p_h1.add_run('一、 學生隨堂練習表（請在空白欄寫出國字或注音）')
r_h1.font.name = '微軟正黑體'
r_h1._element.rPr.rFonts.set(qn('w:eastAsia'), '微軟正黑體')
r_h1.font.size = Pt(12)
r_h1.font.bold = True
r_h1.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)

# Data items (28 items parsed from chinese-pronunciation.docx)
items = [
    (1, '「闔」第光臨', 'ㄏㄜˊ'),
    (2, '晶「ㄧㄥˊ」剔透', '瑩'),
    (3, '呼「ㄏㄨㄢˋ」', '喚'),
    (4, '「椰」子樹', 'ㄧㄝˊ'),
    (5, '「溜」之大吉', 'ㄌㄧㄡ'),
    (6, '「撒」手人寰', 'ㄙㄚ'),
    (7, '「ㄇㄢˊ」荒地帶', '蠻'),
    (8, '「撒」下種子', 'ㄙㄚˇ'),
    (9, '連夜「ㄔㄜˋ」兵', '撤'),
    (10, '「ㄕㄠ」來喜訊', '捎'),
    (11, '四肢痙「ㄌㄨㄢˊ」', '攣'),
    (12, '「襯」托', 'ㄔㄣˋ'),
    (13, '外形「俏」麗', 'ㄑㄧㄠˋ'),
    (14, '「藤」「蔓」', 'ㄊㄥˊ ㄇㄢˋ'),
    (15, '「皎」潔銀白', 'ㄐㄧㄠˇ'),
    (16, '山「峦」起伏', 'ㄌㄨㄢˊ'),
    (17, '熱「忱」', 'ㄔㄣˊ'),
    (18, '「ㄍㄨㄣˇ」動', '滾'),
    (19, '「ㄕㄠ」安勿躁', '稍'),
    (20, '電視「ㄧㄥˊ」幕', '螢'),
    (21, '靠著「ㄌㄢˊ」杆', '欄'),
    (22, '一「ㄇㄟˊ」銅板', '枚'),
    (23, '經「ㄧㄥˊ」不善', '營'),
    (24, '喜上眉「ㄕㄠ」', '梢'),
    (25, '「ㄌㄨㄢˊ」生兄弟', '孿'),
    (26, '和「ㄒㄧㄝˊ」美好', '諧'),
    (27, '無名小「卒」', 'ㄗㄨˊ'),
    (28, '燈光閃「爍」', 'ㄕㄨㄛˋ')
]

# Table 1: Student worksheet table (6 columns)
table1 = doc.add_table(rows=15, cols=6)
table1.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(table1)

headers1 = ['題號', '題目詞語', '學生作答', '題號', '題目詞語', '學生作答']
hdr_cells1 = table1.rows[0].cells
for i, h in enumerate(headers1):
    hdr_cells1[i].text = h
    set_cell_background(hdr_cells1[i], 'D9E1F2')
    set_cell_margins(hdr_cells1[i], top=120, bottom=120, left=100, right=100)
    p = hdr_cells1[i].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.name = '微軟正黑體'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '微軟正黑體')
        run.font.bold = True
        run.font.size = Pt(10.5)

for row_idx in range(14):
    row_cells = table1.rows[row_idx + 1].cells
    
    # Left side (1-14)
    item_l = items[row_idx]
    row_cells[0].text = str(item_l[0])
    row_cells[1].text = item_l[1]
    row_cells[2].text = '' # blank for student answer
    
    # Right side (15-28)
    item_r = items[row_idx + 14]
    row_cells[3].text = str(item_r[0])
    row_cells[4].text = item_r[1]
    row_cells[5].text = '' # blank for student answer
    
    for i, cell in enumerate(row_cells):
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        if i in [0, 2, 3, 5]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.name = '微軟正黑體'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '微軟正黑體')
            run.font.size = Pt(10)

doc.add_paragraph().paragraph_format.space_after = Pt(12)

# Section 2: Teacher Reference Answer Table
p_h2 = doc.add_paragraph()
r_h2 = p_h2.add_run('二、 教師參考對照表（包含目標解答對照）')
r_h2.font.name = '微軟正黑體'
r_h2._element.rPr.rFonts.set(qn('w:eastAsia'), '微軟正黑體')
r_h2.font.size = Pt(12)
r_h2.font.bold = True
r_h2.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)

table2 = doc.add_table(rows=15, cols=6)
table2.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(table2)

headers2 = ['題號', '題目詞語', '參考解答', '題號', '題目詞語', '參考解答']
hdr_cells2 = table2.rows[0].cells
for i, h in enumerate(headers2):
    hdr_cells2[i].text = h
    set_cell_background(hdr_cells2[i], 'F2F2F2')
    set_cell_margins(hdr_cells2[i], top=120, bottom=120, left=100, right=100)
    p = hdr_cells2[i].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.name = '微軟正黑體'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '微軟正黑體')
        run.font.bold = True
        run.font.size = Pt(10.5)

for row_idx in range(14):
    row_cells = table2.rows[row_idx + 1].cells
    
    # Left side
    item_l = items[row_idx]
    row_cells[0].text = str(item_l[0])
    row_cells[1].text = item_l[1]
    row_cells[2].text = item_l[2]
    
    # Right side
    item_r = items[row_idx + 14]
    row_cells[3].text = str(item_r[0])
    row_cells[4].text = item_r[1]
    row_cells[5].text = item_r[2]
    
    for i, cell in enumerate(row_cells):
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        if i in [0, 2, 3, 5]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.name = '微軟正黑體'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '微軟正黑體')
            run.font.size = Pt(10)
            if i in [2, 5]:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00) # Red for answers

output_path1 = 'output/chinese-pronunciation-table.docx'
output_path2 = 'chinese-pronunciation-table.docx'
doc.save(output_path1)
doc.save(output_path2)
print(f'Successfully generated: {output_path1} and {output_path2}')
