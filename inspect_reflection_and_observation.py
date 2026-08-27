import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

def inspect_file(filename):
    print(f"\n==================== {filename} ====================")
    doc = docx.Document(filename)
    print("--- PARAGRAPHS ---")
    for p in doc.paragraphs:
        if p.text.strip():
            print(p.text)
    
    print("\n--- TABLES ---")
    for i, table in enumerate(doc.tables):
        print(f"\nTable {i+1} ({len(table.rows)} rows x {len(table.columns)} cols):")
        for r_idx, row in enumerate(table.rows):
            cells = [c.text.strip().replace('\r', '').replace('\n', ' ') for c in row.cells]
            print(f"  Row {r_idx}: {cells}")

inspect_file('114教學省思心得紀錄表.docx')
inspect_file('114觀課紀錄表.docx')
