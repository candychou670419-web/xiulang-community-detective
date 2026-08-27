import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

print("=== Reading 學習革命的願景讀書會.docx ===")
try:
    doc = docx.Document('學習革命的願景讀書會.docx')
    print(f"Paragraphs count: {len(doc.paragraphs)}")
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            print(f"P{i}: {p.text}")
    print(f"Tables count: {len(doc.tables)}")
    for t_idx, t in enumerate(doc.tables):
        print(f"\n--- Table {t_idx+1} ({len(t.rows)} rows x {len(t.columns)} cols) ---")
        for r_idx, r in enumerate(t.rows):
            print(f"Row {r_idx}: {[c.text.strip().replace('\n', ' ') for c in r.cells]}")
except Exception as e:
    print(f"Docx error: {e}")
