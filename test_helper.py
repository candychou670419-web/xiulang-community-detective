import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import qn, nsdecls
import os

os.makedirs('output', exist_ok=True)

def create_document():
    doc = docx.Document()

    # Set page margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles helper
    def set_run_font(run, name='微軟正黑體', size_pt=10.5, bold=False, color_rgb=(0x33, 0x33, 0x33), italic=False):
        run.font.name = name
        run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
        run.font.size = Pt(size_pt)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = RGBColor(*color_rgb)

    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        set_run_font(run, name='標楷體', size_pt=18, bold=True, color_rgb=(0x1F, 0x4E, 0x79))
        return p

    def add_subtitle(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(18)
        run = p.add_run(text)
        set_run_font(run, name='微軟正黑體', size_pt=12.5, bold=True, color_rgb=(0x2E, 0x75, 0xB6))
        return p

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run(text)
        set_run_font(run, name='微軟正黑體', size_pt=14, bold=True, color_rgb=(0x1F, 0x4E, 0x79))
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        set_run_font(run, name='微軟正黑體', size_pt=12, bold=True, color_rgb=(0x2E, 0x75, 0xB6))
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        set_run_font(run, name='微軟正黑體', size_pt=10.5, bold=True, color_rgb=(0x44, 0x44, 0x44))
        return p

    def add_body(text, indent=True):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.25
        if indent:
            p.paragraph_format.first_line_indent = Pt(21)
        run = p.add_run(text)
        set_run_font(run, name='微軟正黑體', size_pt=10.5, bold=False, color_rgb=(0x22, 0x22, 0x22))
        return p

    def add_quote(text):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.4)
        p.paragraph_format.right_indent = Inches(0.4)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.2
        run = p.add_run(text)
        set_run_font(run, name='楷體_GB2312' if '楷體' in [f.name for f in doc.styles['Normal'].font.__class__.__subclasses__()] else '標楷體', size_pt=9.5, italic=True, color_rgb=(0x44, 0x44, 0x55))
        return p

    return doc, add_title, add_subtitle, add_h1, add_h2, add_h3, add_body, add_quote, set_run_font

print("Document helper defined.")
